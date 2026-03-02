"""
將靈寵美工規格書（assets/pets/README.md）轉成 PPT 簡報
靈犀 App — 24 節氣靈寵美工規格
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── 色彩定義 ──
GOLD = RGBColor(0xE8, 0xC5, 0x47)
DARK_BG = RGBColor(0x08, 0x08, 0x0F)
DARK_CARD = RGBColor(0x14, 0x14, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GOLD = RGBColor(0xF5, 0xE6, 0x9A)
DIM_TEXT = RGBColor(0x88, 0x80, 0x70)
BLUE_ACCENT = RGBColor(0x64, 0xB4, 0xFF)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
GREEN = RGBColor(0x64, 0xC8, 0x78)
ORANGE = RGBColor(0xE8, 0x9C, 0x47)
RED = RGBColor(0xC4, 0x40, 0x40)

# 季節色
SPRING_GREEN = RGBColor(0x7B, 0xC9, 0x6F)
SPRING_PINK = RGBColor(0xF0, 0xA0, 0xC0)
SUMMER_ORANGE = RGBColor(0xFF, 0x6B, 0x35)
SUMMER_RED = RGBColor(0xC4, 0x40, 0x40)
AUTUMN_AMBER = RGBColor(0xD4, 0xA5, 0x74)
AUTUMN_SILVER = RGBColor(0xC0, 0xC0, 0xC0)
WINTER_ICE = RGBColor(0x64, 0xB4, 0xFF)
WINTER_PURPLE = RGBColor(0x4A, 0x20, 0x80)
WINTER_AURORA = RGBColor(0x00, 0xFF, 0xA0)

# 五行色
WOOD_C = RGBColor(0x7B, 0xC9, 0x6F)
FIRE_C = RGBColor(0xFF, 0x6B, 0x35)
EARTH_C = RGBColor(0xFF, 0xC1, 0x07)
METAL_C = RGBColor(0xC0, 0xC0, 0xC0)
WATER_C = RGBColor(0x64, 0xB4, 0xFF)

ELEMENT_COLORS = {"木": WOOD_C, "火": FIRE_C, "土": EARTH_C, "金": METAL_C, "水": WATER_C}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def set_slide_bg(slide, color=DARK_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=14,
                color=WHITE, bold=False, align=PP_ALIGN.LEFT,
                font_name="Microsoft JhengHei"):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_rect(slide, left, top, width, height,
             fill_color=DARK_CARD, border_color=None):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_multi_text(slide, left, top, width, height, lines,
                   default_size=13, default_color=WHITE):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        text = item[0]
        size = item[1] if len(item) > 1 else default_size
        color = item[2] if len(item) > 2 else default_color
        bold = item[3] if len(item) > 3 else False
        align = item[4] if len(item) > 4 else PP_ALIGN.LEFT
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Microsoft JhengHei"
        p.alignment = align
        p.space_before = Pt(1)
        p.space_after = Pt(1)
    return txBox


def add_code_block(slide, left, top, width, height, text, font_size=10):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    fill = txBox.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0x1A, 0x1A, 0x28)
    for i, line in enumerate(text.split("\n")):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = LIGHT_GOLD
        p.font.name = "Consolas"
        p.space_before = Pt(1)
        p.space_after = Pt(1)
    return txBox


# ── 24 隻靈寵資料 ──
PETS = [
    # 春（1-6）
    {"num": 1,  "term": "立春", "date": "2/4–2/18",   "name": "青芽鹿", "creature": "鹿",     "element": "木", "zodiac": "水瓶座", "emoji": "🦌", "keywords": "春芽鹿角、新生嫩綠、破土而出的生機",        "folder": "01-lichun"},
    {"num": 2,  "term": "雨水", "date": "2/19–3/4",   "name": "潤澤蛙", "creature": "蛙",     "element": "水", "zodiac": "雙魚座", "emoji": "🐸", "keywords": "雨滴環繞、荷葉坐騎、水珠光澤",          "folder": "02-yushui"},
    {"num": 3,  "term": "驚蟄", "date": "3/5–3/19",   "name": "雷蟲龍", "creature": "龍蟲",   "element": "木", "zodiac": "雙魚座", "emoji": "🐛", "keywords": "微型龍形、雷電紋路、破土甦醒",          "folder": "03-jingzhe"},
    {"num": 4,  "term": "春分", "date": "3/20–4/3",   "name": "衡翼蝶", "creature": "蝴蝶",   "element": "木", "zodiac": "牡羊座", "emoji": "🦋", "keywords": "日夜雙翼、陰陽平衡、花粉光點",          "folder": "04-chunfen"},
    {"num": 5,  "term": "清明", "date": "4/4–4/19",   "name": "清風鶴", "creature": "鶴",     "element": "木", "zodiac": "牡羊座", "emoji": "🕊️", "keywords": "仙鶴身姿、清澈微風、柳絮飄飛",          "folder": "05-qingming"},
    {"num": 6,  "term": "穀雨", "date": "4/20–5/4",   "name": "穀靈兔", "creature": "兔",     "element": "土", "zodiac": "金牛座", "emoji": "🐰", "keywords": "穀穗花環、春雨滋潤、豐收祈願",          "folder": "06-guyu"},
    # 夏（7-12）
    {"num": 7,  "term": "立夏", "date": "5/5–5/20",   "name": "炎蟬精", "creature": "蟬",     "element": "火", "zodiac": "金牛座", "emoji": "🪲", "keywords": "透明薄翼、初夏暖光、蟬鳴音波",          "folder": "07-lixia"},
    {"num": 8,  "term": "小滿", "date": "5/21–6/4",   "name": "金穗狐", "creature": "狐",     "element": "火", "zodiac": "雙子座", "emoji": "🦊", "keywords": "金色毛皮、麥穗尾巴、豐盈飽滿",          "folder": "08-xiaoman"},
    {"num": 9,  "term": "芒種", "date": "6/5–6/20",   "name": "芒鳳雀", "creature": "雀鳳",   "element": "火", "zodiac": "雙子座", "emoji": "🐦", "keywords": "金色羽翼、稻芒光芒、播種希望",          "folder": "09-mangzhong"},
    {"num": 10, "term": "夏至", "date": "6/21–7/6",   "name": "日輪獅", "creature": "獅",     "element": "火", "zodiac": "巨蟹座", "emoji": "🦁", "keywords": "太陽鬃毛、正午光環、最長白晝",          "folder": "10-xiazhi"},
    {"num": 11, "term": "小暑", "date": "7/7–7/21",   "name": "螢火靈", "creature": "螢火蟲", "element": "火", "zodiac": "巨蟹座", "emoji": "✨", "keywords": "巨大螢火蟲、夏夜微光、溫暖引路",        "folder": "11-xiaoshu"},
    {"num": 12, "term": "大暑", "date": "7/22–8/6",   "name": "烈陽鷹", "creature": "鷹",     "element": "土", "zodiac": "獅子座", "emoji": "🦅", "keywords": "烈日雙翼、熱浪氣場、高空俯瞰",          "folder": "12-dashu"},
    # 秋（13-18）
    {"num": 13, "term": "立秋", "date": "8/7–8/22",   "name": "金風虎", "creature": "虎",     "element": "金", "zodiac": "獅子座", "emoji": "🐅", "keywords": "秋風金紋、落葉環繞、威嚴轉涼",          "folder": "13-liqiu"},
    {"num": 14, "term": "處暑", "date": "8/23–9/6",   "name": "涼蟬仙", "creature": "蟬",     "element": "金", "zodiac": "處女座", "emoji": "🪲", "keywords": "褪色薄翼、秋涼微風、暑氣消散",          "folder": "14-chushu"},
    {"num": 15, "term": "白露", "date": "9/7–9/22",   "name": "露珠蛇", "creature": "蛇",     "element": "金", "zodiac": "處女座", "emoji": "🐍", "keywords": "晶瑩鱗片、晨露凝結、月光折射",          "folder": "15-bailu"},
    {"num": 16, "term": "秋分", "date": "9/23–10/7",  "name": "月衡鶴", "creature": "鶴",     "element": "金", "zodiac": "天秤座", "emoji": "🕊️", "keywords": "月光翅膀、天秤光陣、晝夜等分",          "folder": "16-qiufen"},
    {"num": 17, "term": "寒露", "date": "10/8–10/22", "name": "霜菊貓", "creature": "貓",     "element": "水", "zodiac": "天秤座", "emoji": "🐱", "keywords": "菊花冠飾、寒霜毛色、深秋寧靜",          "folder": "17-hanlu"},
    {"num": 18, "term": "霜降", "date": "10/23–11/6", "name": "霜狼靈", "creature": "狼",     "element": "水", "zodiac": "天蠍座", "emoji": "🐺", "keywords": "銀霜皮毛、冰晶呼吸、孤月嚎叫",          "folder": "18-shuangjiang"},
    # 冬（19-24）
    {"num": 19, "term": "立冬", "date": "11/7–11/21", "name": "冬眠熊", "creature": "熊",     "element": "水", "zodiac": "天蠍座", "emoji": "🐻", "keywords": "厚實毛皮、冬眠蜷縮、蓄能守護",          "folder": "19-lidong"},
    {"num": 20, "term": "小雪", "date": "11/22–12/6", "name": "雪兔仙", "creature": "兔",     "element": "水", "zodiac": "射手座", "emoji": "🐇", "keywords": "白色長耳、初雪飄落、輕盈跳躍",          "folder": "20-xiaoxue"},
    {"num": 21, "term": "大雪", "date": "12/7–12/21", "name": "雪鴞靈", "creature": "貓頭鷹", "element": "水", "zodiac": "射手座", "emoji": "🦉", "keywords": "雪白羽毛、暗夜慧眼、風雪中守望",        "folder": "21-daxue"},
    {"num": 22, "term": "冬至", "date": "12/22–1/4",  "name": "玄冰龍", "creature": "龍",     "element": "水", "zodiac": "摩羯座", "emoji": "🐉", "keywords": "冰晶龍鱗、極寒吐息、冬至轉陽",          "folder": "22-dongzhi"},
    {"num": 23, "term": "小寒", "date": "1/5–1/19",   "name": "寒星鯨", "creature": "鯨",     "element": "水", "zodiac": "摩羯座", "emoji": "🐋", "keywords": "星辰鯨身、深海潛行、寒夜星空",          "folder": "23-xiaohan"},
    {"num": 24, "term": "大寒", "date": "1/20–2/3",   "name": "極光鳳", "creature": "鳳凰",   "element": "土", "zodiac": "水瓶座", "emoji": "🔥", "keywords": "極光羽翼、嚴冬中重生、春之預兆",        "folder": "24-dahan"},
]


def pet_card(slide, x, y, w, h, pet, season_color):
    """畫一張靈寵卡片"""
    ec = ELEMENT_COLORS.get(pet["element"], WHITE)
    add_rect(slide, x, y, w, h, DARK_CARD, season_color)
    # emoji
    add_textbox(slide, x, y + 0.02, w, 0.45, pet["emoji"], 26, WHITE, False, PP_ALIGN.CENTER)
    # 名稱 + 節氣
    add_textbox(slide, x + 0.03, y + 0.48, w - 0.06, 0.25, pet["name"], 13, season_color, True, PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.03, y + 0.72, w - 0.06, 0.2, f"{pet['term']}  {pet['date']}", 8, DIM_TEXT, False, PP_ALIGN.CENTER)
    # 屬性
    add_textbox(slide, x + 0.03, y + 0.92, w - 0.06, 0.2,
                f"{pet['element']}  {pet['creature']}  {pet['zodiac']}", 8, ec, False, PP_ALIGN.CENTER)
    # 關鍵字
    add_textbox(slide, x + 0.03, y + 1.12, w - 0.06, 0.7, pet["keywords"], 7, DIM_TEXT, False, PP_ALIGN.CENTER)
    # 資料夾
    add_textbox(slide, x + 0.03, y + 1.65, w - 0.06, 0.2, pet["folder"], 7, RGBColor(0x55, 0x50, 0x45), False, PP_ALIGN.CENTER)


# ════════════════════════════════════════════════
# Slide 1: 封面
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 2, 1.0, 9, 1.0, "靈犀 LingXi", 54, GOLD, True, PP_ALIGN.CENTER)
add_textbox(slide, 2, 2.2, 9, 0.8, "🐾 24 節氣靈寵美工規格書", 30, LIGHT_GOLD, False, PP_ALIGN.CENTER)
add_textbox(slide, 2, 3.3, 9, 0.6, "交給 AI 美工製作的完整指引", 16, WHITE, False, PP_ALIGN.CENTER)

# 季節色帶
seasons = [
    ("🌸 春", "6 隻", SPRING_GREEN),
    ("☀️ 夏", "6 隻", GOLD),
    ("🍂 秋", "6 隻", AUTUMN_AMBER),
    ("❄️ 冬", "6 隻", WINTER_ICE),
]
for i, (label, count, color) in enumerate(seasons):
    x = 2.5 + i * 2.2
    add_rect(slide, x, 4.5, 2.0, 0.8, DARK_CARD, color)
    add_textbox(slide, x, 4.52, 2.0, 0.4, label, 16, color, True, PP_ALIGN.CENTER)
    add_textbox(slide, x, 4.92, 2.0, 0.3, count, 12, DIM_TEXT, False, PP_ALIGN.CENTER)

add_textbox(slide, 2, 5.7, 9, 0.5, "24 隻靈寵 × 6 張圖 = 144 張 + 72 組動畫", 18, ORANGE, True, PP_ALIGN.CENTER)
add_textbox(slide, 2, 6.3, 9, 0.4, "風格：東方神秘感 × Q版可愛 · 半透明靈體 · 金色粒子閃爍", 13, DIM_TEXT, False, PP_ALIGN.CENTER)

# ════════════════════════════════════════════════
# Slide 2: 設計規範 + 每隻所需圖片
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "一、設計規範 + 每隻所需圖片", 30, GOLD, True)

# 左側：每隻需要的圖
add_rect(slide, 0.5, 1.1, 6.0, 3.5, DARK_CARD, GOLD)
add_textbox(slide, 0.8, 1.2, 5.5, 0.4, "每隻靈寵需要 6 張圖", 16, GOLD, True)

files = [
    ("avatar.png", "512×512", "頭像（對話氣泡、列表）", "🗣️"),
    ("full.png", "1024×1024", "全身圖（靈寵主頁形象區）", "🖼️"),
    ("evo-1.png", "1024×1024", "進化階段 1（初始型態）", "⭐"),
    ("evo-2.png", "1024×1024", "進化階段 2（覺醒型態）", "⭐⭐"),
    ("evo-3.png", "1024×1024", "進化階段 3（終極型態）", "⭐⭐⭐"),
    ("icon.png", "128×128", "小圖示（選單、通知）", "📌"),
]
for i, (fname, size, desc, emoji) in enumerate(files):
    y = 1.7 + i * 0.46
    add_textbox(slide, 0.8, y, 0.4, 0.3, emoji, 12, WHITE)
    add_textbox(slide, 1.2, y, 1.8, 0.3, fname, 11, LIGHT_GOLD, True)
    add_textbox(slide, 3.0, y, 1.0, 0.3, size, 10, DIM_TEXT)
    add_textbox(slide, 4.1, y, 2.2, 0.3, desc, 10, WHITE)

# 右側：風格規範
add_rect(slide, 6.8, 1.1, 6.0, 3.5, DARK_CARD, PURPLE)
add_textbox(slide, 7.1, 1.2, 5.5, 0.4, "風格規範", 16, PURPLE, True)
add_multi_text(slide, 7.1, 1.7, 5.5, 2.8, [
    ("格式", 12, LIGHT_GOLD, True),
    ("PNG，透明背景", 11, WHITE),
    ("", 4),
    ("風格", 12, LIGHT_GOLD, True),
    ("東方神秘感 × Q 版可愛", 11, WHITE),
    ("帶靈氣光暈、半透明靈體", 11, WHITE),
    ("金色粒子閃爍效果", 11, WHITE),
    ("", 4),
    ("深色背景適配", 12, LIGHT_GOLD, True),
    ("APP 背景 #08080F，靈寵需在深色上顯眼", 11, DIM_TEXT),
    ("", 4),
    ("比例", 12, LIGHT_GOLD, True),
    ("Chibi（Q版）比例，可愛親切", 11, WHITE),
    ("友善表情，靈氣飄浮姿態", 11, WHITE),
])

# 季節色系
add_rect(slide, 0.5, 4.9, 12.3, 2.3, DARK_CARD, GOLD)
add_textbox(slide, 0.8, 5.0, 11.8, 0.4, "季節色系對照", 16, GOLD, True)

season_colors = [
    ("🌸 春（01-06）", "嫩綠 · 粉色 · 淺藍", "#7BC96F  #F0A0C0  #A0D8EF",
     [SPRING_GREEN, SPRING_PINK, RGBColor(0xA0, 0xD8, 0xEF)]),
    ("☀️ 夏（07-12）", "金色 · 橙紅 · 烈焰", "#E8C547  #FF6B35  #C44040",
     [GOLD, SUMMER_ORANGE, SUMMER_RED]),
    ("🍂 秋（13-18）", "金黃 · 琥珀 · 銀白", "#D4A574  #C0A030  #C0C0C0",
     [AUTUMN_AMBER, RGBColor(0xC0, 0xA0, 0x30), AUTUMN_SILVER]),
    ("❄️ 冬（19-24）", "冰藍 · 深紫 · 極光", "#64B4FF  #4A2080  #00FFA0",
     [WINTER_ICE, WINTER_PURPLE, WINTER_AURORA]),
]
for i, (label, desc, codes, colors) in enumerate(season_colors):
    x = 0.8 + i * 3.1
    add_textbox(slide, x, 5.45, 2.8, 0.3, label, 12, colors[0], True)
    add_textbox(slide, x, 5.75, 2.8, 0.25, desc, 10, DIM_TEXT)
    # 色塊
    for j, c in enumerate(colors):
        add_rect(slide, x + j * 0.9, 6.05, 0.8, 0.35, c)
    add_textbox(slide, x, 6.45, 2.8, 0.2, codes, 7, DIM_TEXT)

# ════════════════════════════════════════════════
# Slide 3: 春季靈寵（1-6）
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "二、🌸 春季靈寵（立春 → 穀雨）", 30, SPRING_GREEN, True)
add_textbox(slide, 0.5, 0.85, 8, 0.3, "主色調：嫩綠 #7BC96F · 粉色 #F0A0C0 · 淺藍 #A0D8EF", 12, DIM_TEXT)

for i in range(6):
    x = 0.3 + i * 2.1
    pet_card(slide, x, 1.3, 1.95, 1.9, PETS[i], SPRING_GREEN)

# 下方：各自設計關鍵字展開
for i in range(6):
    x = 0.3 + i * 2.1
    pet = PETS[i]
    ec = ELEMENT_COLORS.get(pet["element"], WHITE)
    add_rect(slide, x, 3.4, 1.95, 3.8, RGBColor(0x10, 0x10, 0x18), ec)
    add_textbox(slide, x + 0.03, 3.45, 1.89, 0.25, f"{pet['emoji']} {pet['name']}", 11, ec, True, PP_ALIGN.CENTER)
    add_multi_text(slide, x + 0.05, 3.72, 1.85, 3.3, [
        ("靈獸原型", 8, LIGHT_GOLD, True),
        (pet["creature"], 9, WHITE),
        ("", 3),
        ("五行", 8, LIGHT_GOLD, True),
        (pet["element"], 9, ec),
        ("", 3),
        ("星座", 8, LIGHT_GOLD, True),
        (pet["zodiac"], 9, WHITE),
        ("", 3),
        ("設計關鍵字", 8, LIGHT_GOLD, True),
        (pet["keywords"], 8, DIM_TEXT),
        ("", 3),
        ("資料夾", 8, LIGHT_GOLD, True),
        (pet["folder"] + "/", 8, DIM_TEXT),
    ])

# ════════════════════════════════════════════════
# Slide 4: 夏季靈寵（7-12）
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "三、☀️ 夏季靈寵（立夏 → 大暑）", 30, GOLD, True)
add_textbox(slide, 0.5, 0.85, 8, 0.3, "主色調：金色 #E8C547 · 橙紅 #FF6B35 · 烈焰 #C44040", 12, DIM_TEXT)

for i in range(6):
    x = 0.3 + i * 2.1
    pet_card(slide, x, 1.3, 1.95, 1.9, PETS[6 + i], GOLD)

for i in range(6):
    x = 0.3 + i * 2.1
    pet = PETS[6 + i]
    ec = ELEMENT_COLORS.get(pet["element"], WHITE)
    add_rect(slide, x, 3.4, 1.95, 3.8, RGBColor(0x10, 0x10, 0x18), ec)
    add_textbox(slide, x + 0.03, 3.45, 1.89, 0.25, f"{pet['emoji']} {pet['name']}", 11, ec, True, PP_ALIGN.CENTER)
    add_multi_text(slide, x + 0.05, 3.72, 1.85, 3.3, [
        ("靈獸原型", 8, LIGHT_GOLD, True),
        (pet["creature"], 9, WHITE),
        ("", 3),
        ("五行", 8, LIGHT_GOLD, True),
        (pet["element"], 9, ec),
        ("", 3),
        ("星座", 8, LIGHT_GOLD, True),
        (pet["zodiac"], 9, WHITE),
        ("", 3),
        ("設計關鍵字", 8, LIGHT_GOLD, True),
        (pet["keywords"], 8, DIM_TEXT),
        ("", 3),
        ("資料夾", 8, LIGHT_GOLD, True),
        (pet["folder"] + "/", 8, DIM_TEXT),
    ])

# ════════════════════════════════════════════════
# Slide 5: 秋季靈寵（13-18）
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "四、🍂 秋季靈寵（立秋 → 霜降）", 30, AUTUMN_AMBER, True)
add_textbox(slide, 0.5, 0.85, 8, 0.3, "主色調：金黃 #D4A574 · 琥珀 #C0A030 · 銀白 #C0C0C0", 12, DIM_TEXT)

for i in range(6):
    x = 0.3 + i * 2.1
    pet_card(slide, x, 1.3, 1.95, 1.9, PETS[12 + i], AUTUMN_AMBER)

for i in range(6):
    x = 0.3 + i * 2.1
    pet = PETS[12 + i]
    ec = ELEMENT_COLORS.get(pet["element"], WHITE)
    add_rect(slide, x, 3.4, 1.95, 3.8, RGBColor(0x10, 0x10, 0x18), ec)
    add_textbox(slide, x + 0.03, 3.45, 1.89, 0.25, f"{pet['emoji']} {pet['name']}", 11, ec, True, PP_ALIGN.CENTER)
    add_multi_text(slide, x + 0.05, 3.72, 1.85, 3.3, [
        ("靈獸原型", 8, LIGHT_GOLD, True),
        (pet["creature"], 9, WHITE),
        ("", 3),
        ("五行", 8, LIGHT_GOLD, True),
        (pet["element"], 9, ec),
        ("", 3),
        ("星座", 8, LIGHT_GOLD, True),
        (pet["zodiac"], 9, WHITE),
        ("", 3),
        ("設計關鍵字", 8, LIGHT_GOLD, True),
        (pet["keywords"], 8, DIM_TEXT),
        ("", 3),
        ("資料夾", 8, LIGHT_GOLD, True),
        (pet["folder"] + "/", 8, DIM_TEXT),
    ])

# ════════════════════════════════════════════════
# Slide 6: 冬季靈寵（19-24）
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "五、❄️ 冬季靈寵（立冬 → 大寒）", 30, WINTER_ICE, True)
add_textbox(slide, 0.5, 0.85, 8, 0.3, "主色調：冰藍 #64B4FF · 深紫 #4A2080 · 極光 #00FFA0", 12, DIM_TEXT)

for i in range(6):
    x = 0.3 + i * 2.1
    pet_card(slide, x, 1.3, 1.95, 1.9, PETS[18 + i], WINTER_ICE)

for i in range(6):
    x = 0.3 + i * 2.1
    pet = PETS[18 + i]
    ec = ELEMENT_COLORS.get(pet["element"], WHITE)
    add_rect(slide, x, 3.4, 1.95, 3.8, RGBColor(0x10, 0x10, 0x18), ec)
    add_textbox(slide, x + 0.03, 3.45, 1.89, 0.25, f"{pet['emoji']} {pet['name']}", 11, ec, True, PP_ALIGN.CENTER)
    add_multi_text(slide, x + 0.05, 3.72, 1.85, 3.3, [
        ("靈獸原型", 8, LIGHT_GOLD, True),
        (pet["creature"], 9, WHITE),
        ("", 3),
        ("五行", 8, LIGHT_GOLD, True),
        (pet["element"], 9, ec),
        ("", 3),
        ("星座", 8, LIGHT_GOLD, True),
        (pet["zodiac"], 9, WHITE),
        ("", 3),
        ("設計關鍵字", 8, LIGHT_GOLD, True),
        (pet["keywords"], 8, DIM_TEXT),
        ("", 3),
        ("資料夾", 8, LIGHT_GOLD, True),
        (pet["folder"] + "/", 8, DIM_TEXT),
    ])

# ════════════════════════════════════════════════
# Slide 7: 進化系統
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "六、進化系統（3 階段）", 30, GOLD, True)

# 三階段
stages = [
    ("⭐", "evo-1.png", "初始型態", "small, baby-like\nsimple aura\nsingle color glow\ninnocent expression",
     "體型小、簡單光環\n單色光暈、天真表情\n剛被召喚的幼生體", DIM_TEXT),
    ("⭐⭐", "evo-2.png", "覺醒型態", "medium, detailed\nstronger aura\naccessories/armor\ndual color glow, confident",
     "體型中等、更多細節\n配件/輕甲裝飾\n雙色光暈、自信神態", GOLD),
    ("⭐⭐⭐", "evo-3.png", "終極型態", "majestic, full armor\nintense multi-color aura\nconstellation patterns\ndivine",
     "體型威武、完整裝甲\n多色強烈光環\n星座紋路、神聖感", PURPLE),
]
for i, (stars, fname, label, en_kw, zh_desc, color) in enumerate(stages):
    x = 0.5 + i * 4.2
    add_rect(slide, x, 1.1, 3.9, 5.5, DARK_CARD, color)
    add_textbox(slide, x, 1.2, 3.9, 0.4, stars, 22, GOLD, False, PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.1, 1.6, 3.7, 0.35, label, 18, color, True, PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.1, 1.95, 3.7, 0.25, fname + "  |  1024×1024", 10, DIM_TEXT, False, PP_ALIGN.CENTER)

    # 中文描述
    add_rect(slide, x + 0.15, 2.3, 3.6, 1.5, RGBColor(0x1A, 0x1A, 0x28))
    add_textbox(slide, x + 0.2, 2.32, 3.5, 0.2, "設計描述", 10, LIGHT_GOLD, True)
    add_textbox(slide, x + 0.2, 2.55, 3.5, 1.1, zh_desc, 10, WHITE)

    # 英文 prompt 關鍵字
    add_rect(slide, x + 0.15, 3.95, 3.6, 1.5, RGBColor(0x1A, 0x1A, 0x28))
    add_textbox(slide, x + 0.2, 3.97, 3.5, 0.2, "Prompt 關鍵字（英文）", 10, LIGHT_GOLD, True)
    add_code_block(slide, x + 0.2, 4.2, 3.5, 1.1, en_kw, 9)

    # 範例
    add_textbox(slide, x + 0.1, 5.6, 3.7, 0.7,
        f"範例：青芽鹿 {label}\n🦌 → 🦌{'✨' * (i+1)}", 10, DIM_TEXT, False, PP_ALIGN.CENTER)

# 進化觸發
add_rect(slide, 0.5, 6.8, 12.3, 0.5, DARK_CARD, ORANGE)
add_textbox(slide, 0.8, 6.82, 11.8, 0.4,
    "觸發等級：Lv.10（覺醒）· Lv.20（終極）  |  等級上限：免費 Lv.10 / 會員 Lv.20 / 至尊無限",
    12, ORANGE, False, PP_ALIGN.CENTER)

# ════════════════════════════════════════════════
# Slide 8: 動畫 + AI Prompt 模板
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 10, 0.6, "七、動畫素材 + AI 生圖 Prompt 模板", 30, GOLD, True)

# 動畫
add_rect(slide, 0.5, 1.1, 5.8, 2.5, DARK_CARD, BLUE_ACCENT)
add_textbox(slide, 0.8, 1.2, 5.3, 0.4, "動畫素材（Lottie JSON）", 16, BLUE_ACCENT, True)
add_textbox(slide, 0.8, 1.6, 5.3, 0.25, "assets/pets/animations/", 11, DIM_TEXT)
add_multi_text(slide, 0.8, 1.9, 5.3, 1.5, [
    ("{pet-id}-idle.json", 12, LIGHT_GOLD, True),
    ("待機動畫：輕微浮動、呼吸光效", 10, DIM_TEXT),
    ("{pet-id}-happy.json", 12, LIGHT_GOLD, True),
    ("開心動畫：跳躍、發光、星星特效", 10, DIM_TEXT),
    ("{pet-id}-talk.json", 12, LIGHT_GOLD, True),
    ("說話動畫：嘴巴微動、表情變化", 10, DIM_TEXT),
])

# AI Prompt 通用模板
add_rect(slide, 6.6, 1.1, 6.2, 2.5, DARK_CARD, PURPLE)
add_textbox(slide, 6.9, 1.2, 5.7, 0.4, "AI 生圖通用 Prompt 模板", 16, PURPLE, True)

prompt_template = """A cute mystical [靈獸名稱] spirit pet,
Chinese fantasy style,
[季節色系] glowing aura,
semi-transparent ethereal body,
inspired by "[節氣英文]"
and [星座] constellation,
dark background (#08080f),
golden sparkle particles,
chibi proportions,
friendly expression,
mystical floating pose,
[設計關鍵字],
high quality game art,
PNG transparent background, 1024x1024"""
add_code_block(slide, 6.9, 1.65, 5.7, 1.85, prompt_template, 8)

# 範例 1：青芽鹿
add_rect(slide, 0.5, 3.85, 6.2, 3.4, DARK_CARD, SPRING_GREEN)
add_textbox(slide, 0.8, 3.95, 5.7, 0.4, "範例 Prompt：🦌 青芽鹿（立春·木·水瓶座）", 14, SPRING_GREEN, True)

prompt1 = """A cute mystical young deer spirit pet,
Chinese fantasy style,
fresh green glowing aura with spring buds
growing from its antlers,
inspired by "Beginning of Spring"
and Aquarius constellation,
dark background (#08080f),
golden sparkle particles,
chibi proportions, hopeful expression,
new sprouts and petals floating,
semi-transparent ethereal body,
high quality game art,
PNG transparent background, 1024x1024"""
add_code_block(slide, 0.8, 4.4, 5.7, 2.7, prompt1, 9)

# 範例 2：玄冰龍
add_rect(slide, 6.9, 3.85, 6.0, 3.4, DARK_CARD, WINTER_ICE)
add_textbox(slide, 7.2, 3.95, 5.5, 0.4, "範例 Prompt：🐉 玄冰龍（冬至·水·摩羯座）", 14, WINTER_ICE, True)

prompt2 = """A cute mystical ice dragon spirit pet,
Chinese fantasy style,
icy blue (#64b4ff) glowing aura
with frost crystal scales,
inspired by "Winter Solstice"
and Capricorn constellation,
dark background (#08080f),
aurora borealis particles,
chibi proportions,
serene yet powerful expression,
ice breath effect,
constellation pattern on body,
high quality game art,
PNG transparent background, 1024x1024"""
add_code_block(slide, 7.2, 4.4, 5.5, 2.7, prompt2, 9)

# ════════════════════════════════════════════════
# Slide 9: 資料夾結構 + 總統計
# ════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_textbox(slide, 0.5, 0.3, 8, 0.6, "八、資料夾結構 + 總數量統計", 30, GOLD, True)

# 資料夾
folder_tree = """assets/pets/
├── 01-lichun/        # 立春 — 青芽鹿
│   ├── avatar.png    (512×512)
│   ├── full.png      (1024×1024)
│   ├── evo-1.png     (1024×1024)
│   ├── evo-2.png     (1024×1024)
│   ├── evo-3.png     (1024×1024)
│   └── icon.png      (128×128)
├── 02-yushui/        # 雨水 — 潤澤蛙
├── 03-jingzhe/       # 驚蟄 — 雷蟲龍
│   ...
├── 24-dahan/         # 大寒 — 極光鳳
└── animations/
    ├── 01-lichun-idle.json
    ├── 01-lichun-happy.json
    ├── 01-lichun-talk.json
    │   ...
    └── 24-dahan-talk.json"""
add_code_block(slide, 0.5, 1.0, 6.5, 5.5, folder_tree, 11)

# 統計
add_rect(slide, 7.3, 1.0, 5.5, 4.0, DARK_CARD, ORANGE)
add_textbox(slide, 7.6, 1.1, 5.0, 0.4, "總數量統計", 18, ORANGE, True)

stats = [
    ("靈寵數量", "24 隻", GOLD),
    ("", "", WHITE),
    ("每隻圖片", "6 張", WHITE),
    ("  avatar.png", "×24 = 24 張", DIM_TEXT),
    ("  full.png", "×24 = 24 張", DIM_TEXT),
    ("  evo-1/2/3.png", "×24 = 72 張", DIM_TEXT),
    ("  icon.png", "×24 = 24 張", DIM_TEXT),
    ("", "", WHITE),
    ("每隻動畫", "3 組", WHITE),
    ("  idle/happy/talk", "×24 = 72 組", DIM_TEXT),
]
for i, (label, count, color) in enumerate(stats):
    y = 1.6 + i * 0.22
    add_textbox(slide, 7.6, y, 3.0, 0.2, label, 10, color)
    add_textbox(slide, 10.6, y, 1.8, 0.2, count, 10, color, True, PP_ALIGN.RIGHT)

add_rect(slide, 7.5, 4.15, 5.1, 0.4, ORANGE)
add_textbox(slide, 7.6, 4.17, 4.9, 0.35, "圖片合計 144 張 + 動畫 72 組", 14, DARK_BG, True, PP_ALIGN.CENTER)

# 交付格式
add_rect(slide, 7.3, 4.8, 5.5, 1.7, DARK_CARD, GOLD)
add_textbox(slide, 7.6, 4.9, 5.0, 0.35, "交付格式", 14, GOLD, True)
add_multi_text(slide, 7.6, 5.3, 5.0, 1.1, [
    ("圖片：PNG 透明背景，sRGB", 11, WHITE),
    ("動畫：Lottie JSON 格式", 11, WHITE),
    ("命名：全小寫，連字號分隔", 11, WHITE),
    ("品質：無壓縮，保持清晰度", 11, WHITE),
])

# ════════════════════════════════════════════════
# 儲存
# ════════════════════════════════════════════════
out_dir = r"G:\共用雲端硬碟\有泉科技有限公司\內部開發\APP\算命系統\LingXi"
out_path = os.path.join(out_dir, "LingXi_Pet_Spec.pptx")
prs.save(out_path)
print(f"PPT saved: {out_path}")
